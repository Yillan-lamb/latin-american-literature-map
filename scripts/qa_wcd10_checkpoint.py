#!/usr/bin/env python3
"""Verify WCD-10 actual master, all export formats and public non-expansion."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sqlite3
import subprocess
import tempfile
from contextlib import closing
from pathlib import Path

from export_from_sqlite import table_names, table_rows
from validate_master import validate_database

B01_BASELINE = "05f3372c69a54ec331d31b05d16c8288d6d3a993"
B01_SQL_SHA256 = "2bbb93a2b2bfd3a8ffed9906c5db11d94b08d2d78d363c295aa310658eed0ffc"


def baseline_bytes(repo, baseline_path):
    if baseline_path is not None:
        return baseline_path.read_bytes()
    return subprocess.check_output([
        "git", "show", B01_BASELINE + ":data/master/V1_MASTER.sqlite"], cwd=repo)


def verify_reviewed_delta(repo, tables, mapping, *, baseline_path=None,
                          migration_name="0036_wcd10_b01_author_works.sql", reviewed_sha256=B01_SQL_SHA256,
                          reviewed_reviewer=None):
    """Replay exact reviewed SQL. Defaults stay bound to immutable B01."""
    migration = repo / "data/master/migrations" / migration_name
    if digest(migration) != reviewed_sha256:
        raise ValueError("migration differs from independently reviewed SQL")
    with tempfile.TemporaryDirectory(prefix="wcd10-exact-replay-") as temporary:
        baseline = Path(temporary) / "baseline.sqlite"
        baseline.write_bytes(baseline_bytes(repo, baseline_path))
        if digest(baseline) != mapping["candidate_review"]["baseline_database_sha256"]:
            raise ValueError("historical baseline hash mismatch")
        with closing(sqlite3.connect(baseline)) as conn:
            conn.row_factory = sqlite3.Row
            conn.execute("PRAGMA foreign_keys=ON")
            conn.executescript("BEGIN;\n" + migration.read_text(encoding="utf-8") + "\nCOMMIT;")
            expected_tables = {table: table_rows(conn, table) for table in table_names(conn)}
            if set(tables) != set(expected_tables):
                raise ValueError("unexpected table set after reviewed migration")
            for table, expected in expected_tables.items():
                if table != "migration_log":
                    if tables[table] != expected:
                        raise ValueError("reviewed SQL exact replay mismatch: " + table)
                    continue
                columns, old_rows = expected
                current_columns, current_rows = tables[table]
                new_rows = [row for row in current_rows if row not in old_rows]
                if columns != current_columns or any(row not in current_rows for row in old_rows) or len(new_rows) != 1:
                    raise ValueError("migration log history changed")
                record = new_rows[0]
                fields = {"migration_id": migration.stem, "task_id": mapping["task_id"],
                          "reviewer": (reviewed_reviewer if reviewed_reviewer is not None
                                       else mapping["candidate_review"]["reviewer"]),
                          "sql_sha256": reviewed_sha256, "schema_version": "0.4"}
                if any(record[key] != value for key, value in fields.items()) or not record["applied_at"]:
                    raise ValueError("migration log does not bind approved SQL and reviewer")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify(repo, exports, before_public, after_public, mapping_path, *, baseline_path=None,
           migration_name="0036_wcd10_b01_author_works.sql", reviewed_sha256=B01_SQL_SHA256,
           reviewed_reviewer=None):
    if before_public.resolve() == after_public.resolve():
        raise ValueError("public before and after directories must be distinct")
    mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
    database = repo / "data/master/V1_MASTER.sqlite"
    master = validate_database(database)
    if master["verdict"] != "pass":
        raise ValueError(master["errors"])
    data = json.loads((exports / "data.json").read_text(encoding="utf-8"))
    if data["metadata"]["source_database_sha256"] != digest(database):
        raise ValueError("export belongs to another SQLite input")
    if data["metadata"]["release_status"] != "development_candidate":
        raise ValueError("checkpoint must not be marked released")
    with closing(sqlite3.connect(database.resolve().as_uri() + "?mode=ro", uri=True)) as conn:
        conn.row_factory = sqlite3.Row
        tables = {table: table_rows(conn, table) for table in table_names(conn)}
    if set(data["tables"]) != set(tables):
        raise ValueError("JSON table set differs from SQLite")
    for table, (columns, rows) in tables.items():
        if data["columns"][table] != columns or data["tables"][table] != rows:
            raise ValueError("JSON mismatch: " + table)
        with (exports / (table + ".csv")).open(encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle)
            csv_rows = list(reader)
        expected_csv = [{key: "" if value is None else str(value) for key, value in row.items()} for row in rows]
        if reader.fieldnames != columns or csv_rows != expected_csv:
            raise ValueError("CSV mismatch: " + table)
    from openpyxl import load_workbook
    xlsx = list(exports.glob("*.xlsx"))
    if len(xlsx) != 1:
        raise ValueError("exactly one XLSX required")
    workbook = load_workbook(xlsx[0], read_only=True, data_only=False)
    try:
        if set(workbook.sheetnames) != {table[:31] for table in tables}:
            raise ValueError("XLSX sheet set differs from SQLite")
        for table, (columns, rows) in tables.items():
            sheet = workbook[table[:31]]
            cells = list(sheet.iter_rows())
            if any(cell.data_type == "f" for row in cells for cell in row):
                raise ValueError("unexpected formula in export: " + table)
            values = [["" if cell.value is None else str(cell.value) for cell in row] for row in cells]
            expected = [columns] + [["" if row[key] is None else str(row[key]) for key in columns] for row in rows]
            if values != expected:
                raise ValueError("XLSX mismatch: " + table)
    finally:
        workbook.close()
    verify_reviewed_delta(repo, tables, mapping, baseline_path=baseline_path,
                          migration_name=migration_name, reviewed_sha256=reviewed_sha256,
                          reviewed_reviewer=reviewed_reviewer)
    with tempfile.TemporaryDirectory(prefix="wcd10-preservation-") as temporary:
        baseline = Path(temporary) / "baseline.sqlite"
        blob = baseline_bytes(repo, baseline_path)
        baseline.write_bytes(blob)
        if digest(baseline) != mapping["candidate_review"]["baseline_database_sha256"]:
            raise ValueError("historical baseline hash mismatch")
        with closing(sqlite3.connect(baseline.resolve().as_uri() + "?mode=ro", uri=True)) as conn:
            conn.row_factory = sqlite3.Row
            for table in table_names(conn):
                _, rows = table_rows(conn, table)
                if table == "metadata":
                    old_hash = next(row["value"] for row in rows if row["key"] == "sha256")
                    current_hash = next(row["value"] for row in tables[table][1] if row["key"] == "sha256")
                    if old_hash != current_hash:
                        raise ValueError("historical metadata hash was overwritten")
                    continue
                before = {json.dumps(row, sort_keys=True, ensure_ascii=False) for row in rows}
                after = {json.dumps(row, sort_keys=True, ensure_ascii=False) for row in tables[table][1]}
                if not before <= after:
                    raise ValueError("old rows changed or removed: " + table)
                if table in mapping["counts"] and len(after) != len(before) + mapping["counts"][table]:
                    raise ValueError("delta count mismatch: " + table)
    before = {str(path.relative_to(before_public)): digest(path) for path in before_public.rglob("*") if path.is_file()}
    after = {str(path.relative_to(after_public)): digest(path) for path in after_public.rglob("*") if path.is_file()}
    if before != after:
        changed = sorted(set(before) ^ set(after) | {key for key in before.keys() & after.keys() if before[key] != after[key]})
        raise ValueError("public artifact changed without curation approval: " + repr(changed))
    public = json.loads((after_public / "data/v2/web/site_data.json").read_text(encoding="utf-8"))
    public_ids = {item["target_id"] for item in public["search_index"]}
    added_entities = {value for key, value in mapping["id_map"].items() if "-W" in key and value.startswith("V1-ENT-")}
    if added_entities & public_ids:
        raise ValueError("new Research objects leaked into public routes")
    return {"status": "PASS", "database_sha256": digest(database), "master": master,
            "export_table_formats_equal": {"json": len(tables), "csv": len(tables), "xlsx": len(tables)},
            "existing_rows_preserved": True, "reviewed_sql_exact_replay": True,
            "public_artifacts_byte_identical": len(before),
            "public_entities": len(public_ids), "new_entities_public": 0,
            "snapshot_boundary": "Research checkpoint only; not a new formal Data release"}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--exports", type=Path, required=True)
    parser.add_argument("--before-public", type=Path, required=True)
    parser.add_argument("--after-public", type=Path, required=True)
    parser.add_argument("--mapping", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    report = verify(args.repo, args.exports, args.before_public, args.after_public, args.mapping)
    content = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(content, encoding="utf-8")
    print(content)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
