#!/usr/bin/env python3
"""Verify TASK-127's explicit type correction, exports and public preservation."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sqlite3
from pathlib import Path

from export_from_sqlite import table_names, table_rows
from validate_master import validate_database

ROOT = Path(__file__).resolve().parents[1]
CANDIDATES = ROOT / 'data/changesets/WCD-10/TYPE_CORRECTION/CANDIDATES.json'
MIGRATION = ROOT / 'data/master/migrations/0058_wcd10_collection_type_correction.sql'
EXPECTED_IDS = {'V1-ENT-0460', 'V1-ENT-0465', 'V1-ENT-0489', 'V1-ENT-0494'} | {
    f'V1-ENT-{n:04d}' for n in range(573, 584)
}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def snapshot(path):
    with sqlite3.connect(Path(path).resolve().as_uri() + '?mode=ro', uri=True) as conn:
        conn.row_factory = sqlite3.Row
        return {t: table_rows(conn, t) for t in table_names(conn)}


def verify_delta(before, after, candidate_path=CANDIDATES, migration=MIGRATION):
    candidate = json.loads(Path(candidate_path).read_text())
    if digest(before) != candidate['baseline_sha256']:
        raise ValueError('baseline hash drift')
    if {r['entity_before']['entity_id'] for r in candidate['rows']} != EXPECTED_IDS or len(candidate['rows']) != 15:
        raise ValueError('correction must bind exactly the independently reviewed15 objects')
    a, b = snapshot(before), snapshot(after)
    if set(a) != set(b) or len(a) != 19:
        raise ValueError('table set changed')
    expected_entities = {}
    expected_cards = {}
    for row in candidate['rows']:
        old, new = row['entity_before'], row['entity_after']
        if new != dict(old, entity_type='collection') or old['entity_type'] != 'work':
            raise ValueError('unauthorized entity field change')
        expected_entities[old['entity_id']] = (old, new)
        for old, new in zip(row['cards_before'], row['cards_after'], strict=True):
            expected = dict(old, card_type='collection', content_markdown=old['content_markdown'].replace('- 对象类型：work', '- 对象类型：collection'))
            if old['card_type'] != 'work' or old['content_markdown'].count('- 对象类型：work') != 1 or new != expected:
                raise ValueError('unauthorized card content change')
            expected_cards[old['card_id']] = (old, new)
    if len(expected_cards) != 15:
        raise ValueError('unexpected card count')
    changes = {}
    for table in sorted(a):
        ac, ar = a[table]
        bc, br = b[table]
        if ac != bc:
            raise ValueError('column change: ' + table)
        if table == 'migration_log':
            additional = [r for r in br if r not in ar]
            if len(additional) != 1 or any(r not in br for r in ar):
                raise ValueError('migration history changed')
            r = additional[0]
            if any(r[k] != v for k, v in {'migration_id': Path(migration).stem, 'task_id': 'TASK-127', 'sql_sha256': digest(migration), 'schema_version': '0.4'}.items()):
                raise ValueError('unreviewed migration log')
            changes[table] = 1
            continue
        expected = [dict(r) for r in ar]
        if table in {'entities', 'content_cards'}:
            key = 'entity_id' if table == 'entities' else 'card_id'
            patches = expected_entities if table == 'entities' else expected_cards
            found = set()
            for i, r in enumerate(expected):
                if r[key] in patches:
                    old, new = patches[r[key]]
                    if r != old:
                        raise ValueError('candidate pre-state drift')
                    expected[i] = new
                    found.add(r[key])
            if found != set(patches):
                raise ValueError('missing correction target')
            changes[table] = len(found)
        elif table == 'metadata':
            for r in expected:
                if r['key'] == 'last_change_set':
                    if r['value'] != 'WCD-10-B05':
                        raise ValueError('metadata pre-state drift')
                    r['value'] = 'WCD-10-TYPE-CORRECTION'
            changes[table] = 1
        if expected != br:
            raise ValueError('unapproved row/field difference: ' + table)
    return {'status': 'PASS', 'tables_compared': 19, 'exact_changes': changes, 'unrelated_rows_preserved': True}


def verify_exports(database, exports):
    tables = snapshot(database)
    data = json.loads((exports / 'data.json').read_text())
    if data['metadata']['source_database_sha256'] != digest(database) or set(data['tables']) != set(tables):
        raise ValueError('export input or table set mismatch')
    from openpyxl import load_workbook
    books = list(exports.glob('*.xlsx'))
    if len(books) != 1:
        raise ValueError('exactly one XLSX required')
    workbook = load_workbook(books[0], read_only=True, data_only=False)
    try:
        if set(workbook.sheetnames) != {t[:31] for t in tables}:
            raise ValueError('XLSX table set mismatch')
        for t, (cols, rows) in tables.items():
            if data['columns'][t] != cols or data['tables'][t] != rows:
                raise ValueError('JSON mismatch: ' + t)
            with (exports / (t + '.csv')).open(newline='', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                actual = list(reader)
                if reader.fieldnames != cols or actual != [{k: '' if v is None else str(v) for k, v in r.items()} for r in rows]:
                    raise ValueError('CSV mismatch: ' + t)
            cells = list(workbook[t[:31]].iter_rows())
            if any(c.data_type == 'f' for r in cells for c in r):
                raise ValueError('unexpected Excel formula')
            values = [['' if c.value is None else str(c.value) for c in r] for r in cells]
            if values != [cols] + [['' if r[k] is None else str(r[k]) for k in cols] for r in rows]:
                raise ValueError('XLSX mismatch: ' + t)
    finally:
        workbook.close()
    return {'json': len(tables), 'csv': len(tables), 'xlsx': len(tables)}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--baseline', type=Path, required=True)
    p.add_argument('--database', type=Path, default=ROOT / 'data/master/V1_MASTER.sqlite')
    p.add_argument('--exports', type=Path, required=True)
    p.add_argument('--before-public', type=Path, required=True)
    p.add_argument('--after-public', type=Path, required=True)
    a = p.parse_args()
    master = validate_database(a.database)
    if master['verdict'] != 'pass':
        raise ValueError(master['errors'])
    result = verify_delta(a.baseline, a.database)
    result['export_formats'] = verify_exports(a.database, a.exports)
    def files(folder):
        return {str(x.relative_to(folder)): digest(x) for x in folder.rglob('*') if x.is_file()}
    if a.before_public.resolve() == a.after_public.resolve():
        raise ValueError('public directories must be distinct')
    before, after = files(a.before_public), files(a.after_public)
    if len(before) != 177 or before != after:
        raise ValueError('public payload changed')
    result.update(database_sha256=digest(a.database), public_files_byte_equal=177, master=master)
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))


if __name__ == '__main__':
    main()
